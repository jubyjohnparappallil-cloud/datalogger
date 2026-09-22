"""Parse Elitech / MU-01 temperature-humidity logger PDFs and Excel files."""

from __future__ import annotations

import math
import re
from array import array
from datetime import date, datetime, timedelta
from html import unescape
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pymupdf

EPOCH = datetime(2000, 1, 1)
EPOCH_ORD = date(2000, 1, 1).toordinal()
TEXT_FLAGS = getattr(pymupdf, "TEXTFLAGS_TEXT", 195)

ROW_RE = re.compile(
    r"(\d{2}[-/.]\d{2}[-/.]\d{4})\s+"
    r"(\d{2}:\d{2}(?::\d{2})?)\s+"
    r"(-?\d+\.?\d*)\s+"
    r"(\d+\.?\d*)"
)
# Linux/MuPDF sometimes drops spaces: 21-04-202611:00:0024.051.7
ROW_RE_COMPACT = re.compile(
    r"(\d{2}[-/.]\d{2}[-/.]\d{4})"
    r"(\d{2}:\d{2}(?::\d{2})?)"
    r"(-?\d+\.\d+)"
    r"(\d+\.\d+)"
)

DL_NAME_RE = re.compile(r"(?i)\bDL[\s._-]*(\d+)\b")

DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y")


def logger_column_name(path: Path) -> str:
    stem = path.stem
    match = DL_NAME_RE.search(stem)
    if match:
        return f"DL-{int(match.group(1))}"
    return stem


def logger_sort_key(name: str) -> tuple:
    match = re.search(r"(\d+)$", name)
    if match:
        return (0, int(match.group(1)), name)
    return (1, 0, name.lower())


def _minutes_since_epoch(date_s: str, time_s: str) -> int | None:
    """DD-MM-YYYY / DD/MM/YYYY / YYYY-MM-DD plus HH:MM, no strptime."""
    try:
        if date_s[2] in "-/.":
            first, second, year = int(date_s[0:2]), int(date_s[3:5]), int(date_s[6:10])
            day, month = (second, first) if second > 12 and first <= 12 else (first, second)
        else:
            year, month, day = int(date_s[0:4]), int(date_s[5:7]), int(date_s[8:10])
        hour, minute = int(time_s[0:2]), int(time_s[3:5])
        return (date(year, month, day).toordinal() - EPOCH_ORD) * 1440 + hour * 60 + minute
    except (ValueError, IndexError):
        return None


def _parse_datetime(date_s: str, time_s: str) -> datetime | None:
    minutes = _minutes_since_epoch(date_s, time_s)
    if minutes is None:
        return None
    return EPOCH + timedelta(minutes=minutes)


def _append_text_rows(text: str, minutes: array, temps: array, hums: array) -> None:
    matches = ROW_RE.findall(text)
    if not matches:
        matches = ROW_RE_COMPACT.findall(re.sub(r"\s+", "", text))
    for date_s, time_s, temp_s, rh_s in matches:
        minute = _minutes_since_epoch(date_s, time_s)
        if minute is None:
            continue
        minutes.append(minute)
        temps.append(float(temp_s))
        hums.append(float(rh_s))


def parse_pdf_columns(path: Path) -> tuple[array, array, array]:
    minutes: array = array("i")
    temps: array = array("f")
    hums: array = array("f")
    doc = pymupdf.open(path)
    try:
        for page in doc:
            before = len(minutes)
            _append_text_rows(page.get_text("text") or "", minutes, temps, hums)
            if len(minutes) == before:
                _append_text_rows(page.get_text("text", flags=TEXT_FLAGS) or "", minutes, temps, hums)
    finally:
        doc.close()
    return minutes, temps, hums


def parse_pdf(path: Path) -> list[tuple[datetime, float, float]]:
    minutes, temps, hums = parse_pdf_columns(path)
    return [(EPOCH + timedelta(minutes=m), t, h) for m, t, h in zip(minutes, temps, hums)]


_INLINE_CELL_RE = re.compile(
    r'<c r="([A-Z]+)(\d+)"[^>]*>\s*<is>\s*<t[^>]*>(.*?)</t>',
    re.I | re.S,
)
_VALUE_CELL_RE = re.compile(
    r'<c r="([A-Z]+)(\d+)"[^>]*>\s*<v>(.*?)</v>',
    re.I | re.S,
)


def parse_excel(path: Path) -> list[tuple[datetime, float, float]]:
    data = path.read_bytes()
    if data[:2] == b"PK":
        records = _parse_elitech_xlsx(data)
        if records:
            return records
    return _parse_excel_openpyxl(data)


def _parse_elitech_xlsx(data: bytes) -> list[tuple[datetime, float, float]]:
    """Read Elitech xlsx-in-.xls files from sheet XML. Avoids openpyxl dimension bugs."""
    best: list[tuple[datetime, float, float]] = []
    with ZipFile(BytesIO(data)) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
                continue
            xml = archive.read(name).decode("utf-8", errors="replace").replace("\ufeff", "")
            records = _records_from_sheet_xml(xml)
            if len(records) > len(best):
                best = records
    return best


def _records_from_sheet_xml(xml: str) -> list[tuple[datetime, float, float]]:
    by_row: dict[int, dict[str, str]] = {}
    for col, row, text in _INLINE_CELL_RE.findall(xml):
        by_row.setdefault(int(row), {})[col] = unescape(text)
    if not by_row:
        for col, row, text in _VALUE_CELL_RE.findall(xml):
            by_row.setdefault(int(row), {})[col] = unescape(text)
    if not by_row:
        return []

    header_row = min(by_row)
    header = by_row[header_row]
    cols = sorted(header)
    time_col = temp_col = rh_col = None
    for col in cols:
        label = header[col].strip().lower()
        if time_col is None and ("time" in label or "date" in label):
            time_col = col
        elif temp_col is None and ("temp" in label or "°c" in label or "&#176;" in label or label in {"c", "℃"}):
            temp_col = col
        elif rh_col is None and ("humid" in label or "%rh" in label or label.endswith("rh")):
            rh_col = col
    if time_col is None:
        time_col = cols[1] if len(cols) > 1 else "B"
    if temp_col is None:
        temp_col = cols[2] if len(cols) > 2 else "C"
    if rh_col is None:
        rh_col = cols[3] if len(cols) > 3 else "D"

    records: list[tuple[datetime, float, float]] = []
    for row_no in sorted(by_row):
        if row_no == header_row:
            continue
        cells = by_row[row_no]
        dt = _coerce_excel_datetime(cells.get(time_col))
        if dt is None:
            continue
        try:
            temp = float(cells[temp_col])
            rh = float(cells[rh_col])
        except (KeyError, TypeError, ValueError):
            continue
        records.append((dt, temp, rh))
    return records


def _parse_excel_openpyxl(data: bytes) -> list[tuple[datetime, float, float]]:
    import warnings
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    try:
        sheets = [wb[name] for name in wb.sheetnames if name.lower() in {"list", "data", "readings", "log"}]
        if not sheets:
            sheets = [wb[name] for name in wb.sheetnames]
        best: list[tuple[datetime, float, float]] = []
        for sheet in sheets:
            records = _records_from_openpyxl_sheet(sheet)
            if len(records) > len(best):
                best = records
        return best
    finally:
        wb.close()


def _records_from_openpyxl_sheet(sheet) -> list[tuple[datetime, float, float]]:
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    time_i = temp_i = rh_i = None
    for i, cell in enumerate(header):
        label = str(cell or "").strip().lower()
        if time_i is None and ("time" in label or "date" in label):
            time_i = i
        elif temp_i is None and ("temp" in label or "°c" in label or label in {"c", "℃"}):
            temp_i = i
        elif rh_i is None and ("humid" in label or "%rh" in label or "rh" in label):
            rh_i = i
    if time_i is None:
        time_i = 1
    if temp_i is None:
        temp_i = 2
    if rh_i is None:
        rh_i = 3
    records: list[tuple[datetime, float, float]] = []
    for row in rows:
        if not row or time_i >= len(row):
            continue
        dt = _coerce_excel_datetime(row[time_i])
        if dt is None:
            continue
        try:
            temp = float(row[temp_i])
            rh = float(row[rh_i]) if rh_i < len(row) and row[rh_i] not in (None, "") else None
        except (TypeError, ValueError, IndexError):
            continue
        if rh is None:
            continue
        records.append((dt, temp, rh))
    return records


def _coerce_excel_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    text = str(value).strip()
    match = re.match(r"(\d{2}[-/.]\d{2}[-/.]\d{4})\s+(\d{2}:\d{2}(?::\d{2})?)", text)
    if match:
        return _parse_datetime(match.group(1), match.group(2))
    return None


def parse_logger_file(path: Path) -> list[tuple[datetime, float, float]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix in {".xls", ".xlsx"}:
        return parse_excel(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def parse_logger_columns(path: Path) -> tuple[array, array, array]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf_columns(path)
    if suffix in {".xls", ".xlsx"}:
        return records_to_columns(parse_excel(path))
    raise ValueError(f"Unsupported file type: {path.suffix}")


def sibling_logger_file(path: Path) -> Path | None:
    """If one format comes back empty, try the logger's other file."""
    if path.suffix.lower() == ".pdf":
        for ext in (".xls", ".xlsx"):
            excel = path.with_suffix(ext)
            if excel.exists():
                return excel
        return None
    pdf = path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def parse_one_worker(path_str: str):
    """Parse one logger file into compact arrays."""
    path = Path(path_str)
    logger = logger_column_name(path)
    if not DL_NAME_RE.search(path.stem):
        return logger, None, f"{path.name}: skipped (not a DL logger file)"
    try:
        if not path.is_file():
            return logger, None, f"{path.name}: file missing on server"
        size = path.stat().st_size
        if size < 200:
            return logger, None, f"{path.name}: upload was empty ({size} bytes)"
        columns = parse_logger_columns(path)
        if columns[0]:
            return logger, columns, None
        sibling = sibling_logger_file(path)
        if sibling is not None:
            columns = parse_logger_columns(sibling)
            if columns[0]:
                return logger, columns, None
        return logger, None, f"{path.name}: no temperature/humidity readings found"
    except Exception as exc:
        sibling = sibling_logger_file(path)
        if sibling is not None:
            try:
                columns = parse_logger_columns(sibling)
                if columns[0]:
                    return logger, columns, None
            except Exception:
                pass
        return logger, None, f"{path.name}: {exc}"


def records_to_columns(records: list[tuple[datetime, float, float]]):
    """Pack one logger's readings into compact numeric arrays.

    Storing minutes-since-2000 plus two float arrays instead of Python tuples
    keeps a full 270-logger batch inside a few hundred MB.
    """
    minutes = array("i")
    temps = array("f")
    hums = array("f")
    for dt, temp, rh in records:
        delta = dt.replace(second=0, microsecond=0) - EPOCH
        minutes.append(int(delta.total_seconds()) // 60)
        temps.append(temp)
        hums.append(rh)
    return minutes, temps, hums


def merge_columns(file_columns: dict[str, tuple[array, array, array]]):
    """Align every logger onto one shared timeline, one float column per logger."""
    loggers = sorted(file_columns.keys(), key=logger_sort_key)

    all_minutes: set[int] = set()
    for minutes, _, _ in file_columns.values():
        all_minutes.update(minutes)
    timeline = sorted(all_minutes)
    row_of = {minute: i for i, minute in enumerate(timeline)}
    n_rows = len(timeline)

    temp_cols: dict[str, array] = {}
    hum_cols: dict[str, array] = {}
    for logger in loggers:
        minutes, temps, hums = file_columns[logger]
        temp_col = array("f", [math.nan]) * n_rows
        hum_col = array("f", [math.nan]) * n_rows
        for i, minute in enumerate(minutes):
            row = row_of[minute]
            temp_col[row] = temps[i]
            hum_col[row] = hums[i]
        temp_cols[logger] = temp_col
        hum_cols[logger] = hum_col

    timestamps = [EPOCH + timedelta(minutes=m) for m in timeline]
    return loggers, timestamps, temp_cols, hum_cols
