"""Diagnostic: how many readings actually landed in each DL column."""

import statistics
import sys
import warnings

warnings.filterwarnings("ignore")

from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\Juby John\Desktop\Temperature_Humidity_all_DLs.xlsx"
wb = load_workbook(path, read_only=True, data_only=True)
sheet = wb["Temperature"]
rows = sheet.iter_rows(values_only=True)
header = next(rows)
n = len(header)
filled = [0] * n
total = 0
for row in rows:
    total += 1
    for i in range(2, min(n, len(row))):
        if row[i] is not None:
            filled[i] += 1
wb.close()

print("data rows:", total, "columns:", n)
empty = [header[i] for i in range(2, n) if filled[i] == 0]
print("completely EMPTY columns:", len(empty))
print(empty[:40])
pct = [filled[i] / total * 100 for i in range(2, n)]
print("avg fill %:", round(statistics.mean(pct), 1))
print("min fill %:", round(min(pct), 1), " max fill %:", round(max(pct), 1))
low = sorted(((filled[i] / total * 100, header[i]) for i in range(2, n)))[:15]
print("lowest columns:", [(name, round(p, 1)) for p, name in low])
for i in range(2, 10):
    print(header[i], filled[i], round(filled[i] / total * 100, 1), "%")
